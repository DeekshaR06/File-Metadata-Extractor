import json
import base64
import boto3
import io
import os
from datetime import datetime

s3 = boto3.client('s3')
BUCKET_NAME = os.environ.get('BUCKET_NAME', 'metadata-extractor-temp-deeksha081382')
MAX_FILE_SIZE_BYTES = 6 * 1024 * 1024  # 6 MB


def lambda_handler(event, context):
    try:
        # API Gateway (proxy integration) sends the HTTP body as a string in event['body']
        body = json.loads(event.get('body', '{}'))
        filename = body.get('filename')
        filedata_b64 = body.get('filedata')

        if not filename or not filedata_b64:
            return _response(400, {'error': 'filename and filedata are required'})

        # Decode base64 file data back into raw bytes
        try:
            file_bytes = base64.b64decode(filedata_b64)
        except Exception:
            return _response(400, {'error': 'filedata is not valid base64'})

        # Enforce size limit (defense in depth -- frontend should also check this)
        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            return _response(413, {'error': 'File exceeds 6 MB limit'})

        extension = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''

        # --- Temporary S3 storage step (write, then delete immediately) ---
        s3_key = f"temp-uploads/{context.aws_request_id}-{filename}"
        s3.put_object(Bucket=BUCKET_NAME, Key=s3_key, Body=file_bytes)

        try:
            metadata = _extract_metadata(file_bytes, extension, filename)
        finally:
            # Always delete, even if extraction fails, so nothing lingers
            s3.delete_object(Bucket=BUCKET_NAME, Key=s3_key)

        if metadata is None:
            return _response(415, {'error': f'Unsupported file type: .{extension}'})

        return _response(200, {'filename': filename, 'metadata': metadata})

    except Exception as e:
        return _response(500, {'error': f'Internal error: {str(e)}'})


def _extract_metadata(file_bytes, extension, filename):
    size_bytes = len(file_bytes)

    if extension == 'pdf':
        return _extract_pdf(file_bytes, size_bytes)
    elif extension == 'docx':
        return _extract_docx(file_bytes, size_bytes)
    elif extension in ('jpg', 'jpeg', 'png'):
        return _extract_image(file_bytes, size_bytes)
    elif extension == 'txt':
        return _extract_txt(file_bytes, size_bytes)
    else:
        return None


def _extract_pdf(file_bytes, size_bytes):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    info = reader.metadata or {}

    return {
        'file_type': 'PDF',
        'size_bytes': size_bytes,
        'page_count': len(reader.pages),
        'author': info.get('/Author', 'Unknown'),
        'title': info.get('/Title', 'Unknown'),
        'creation_date': str(info.get('/CreationDate', 'Unknown')),
        'producer': info.get('/Producer', 'Unknown'),
    }


def _extract_docx(file_bytes, size_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    props = doc.core_properties

    word_count = sum(len(p.text.split()) for p in doc.paragraphs)

    return {
        'file_type': 'DOCX',
        'size_bytes': size_bytes,
        'author': props.author or 'Unknown',
        'title': props.title or 'Unknown',
        'created': str(props.created) if props.created else 'Unknown',
        'last_modified_by': props.last_modified_by or 'Unknown',
        'approx_word_count': word_count,
        'paragraph_count': len(doc.paragraphs),
    }


def _extract_image(file_bytes, size_bytes):
    from PIL import Image
    from PIL.ExifTags import TAGS

    img = Image.open(io.BytesIO(file_bytes))
    exif_data = {}

    raw_exif = img.getexif()
    if raw_exif:
        for tag_id, value in raw_exif.items():
            tag_name = TAGS.get(tag_id, tag_id)
            # Keep only simple, JSON-serializable values
            if isinstance(value, (str, int, float)):
                exif_data[str(tag_name)] = value

    return {
        'file_type': img.format,
        'size_bytes': size_bytes,
        'width': img.width,
        'height': img.height,
        'color_mode': img.mode,
        'exif': exif_data if exif_data else 'No EXIF data found',
    }


def _extract_txt(file_bytes, size_bytes):
    try:
        text = file_bytes.decode('utf-8')
        encoding = 'utf-8'
    except UnicodeDecodeError:
        text = file_bytes.decode('latin-1')
        encoding = 'latin-1 (fallback)'

    return {
        'file_type': 'TXT',
        'size_bytes': size_bytes,
        'encoding': encoding,
        'character_count': len(text),
        'line_count': text.count('\n') + 1,
        'word_count': len(text.split()),
    }


def _response(status_code, body_dict):
    return {
        'statusCode': status_code,
        'headers': {
            'Content-Type': 'application/json',
            'Access-Control-Allow-Origin': '*',
        },
        'body': json.dumps(body_dict),
    }